from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl


def inspect_op(inputs: dict, store_):
    """
    Policy:
    - inspect = facts extraction (observation), not planning.
    - Output is control artifacts with stable schema; keep it usable across runs.
    - No file I/O yet; contract-first stub. Later: read/scan/fingerprint/schema inference.
    - Writes to fixed control keys to stabilize downstream integration early.
      (May evolve to keyed-by-source when multi-source is introduced.)
    """

    # inputs: {"source_uri": "...", "read_spec": {...}} を想定
    source_uri = inputs.get("source_uri")
    read_spec = inputs.get("read_spec") or {}

    # 今日の段階では"読む"はしない。契約だけ固定する
    result = {
        "source_uri": source_uri,
        "warnings": [],
        "suggested_read_spec": read_spec,
        "columns": [],  # 将来: [{"name": "...", "dtype": "..."}]
    }

    # runtimeに任せて outputs を永続化する
    return {
        "inspect_result": result,
        "read_spec": result["suggested_read_spec"],
    }


def inspect_filename_kind_op(inputs: dict[str, Any], store_) -> dict[str, Any]:
    """
    Inspect step: Identifies input "kind" from filename pattern only.

    Inputs:
    - input_profile_name: str (profile name to load kind_rules from)
    - path: str (file path to inspect)

    Returns:
    - {"result": {...}} where result is a dict with:
      - schema_version: "inspect_result_v1"
      - input_profile_name: str
      - resolved_path: str
      - filename: str
      - detected_kind: str | None
      - evidence: dict with matched_pattern, matcher
    """
    input_profile_name = inputs.get("input_profile_name")
    if not input_profile_name:
        raise ValueError("input_profile_name is required")

    path_str = inputs.get("path")
    if not path_str:
        raise ValueError("path is required")

    # Load profile config to get kind_rules
    try:
        config = store_.configs.get_spec("input_profile", input_profile_name)
    except KeyError as e:
        raise ValueError(f"input_profile '{input_profile_name}' not found in configs") from e

    kind_rules = config.get("kind_rules")
    if kind_rules is None:
        raise ValueError(f"input_profile '{input_profile_name}' missing 'kind_rules' key")

    # Extract basename from path
    path = Path(path_str)
    filename = path.name

    # Match filename against rules (in order)
    detected_kind = None
    matched_pattern = None

    for rule in kind_rules:
        pattern = rule.get("pattern")
        kind = rule.get("kind")
        if pattern and kind:
            if re.match(pattern, filename):
                detected_kind = kind
                matched_pattern = pattern
                break

    # Build result
    result = {
        "schema_version": "inspect_result_v1",
        "input_profile_name": input_profile_name,
        "resolved_path": path_str,
        "filename": filename,
        "detected_kind": detected_kind,
        "evidence": {
            "matched_pattern": matched_pattern,
            "matcher": "filename_regex",
        },
    }

    return {"result": result}


def _normalize_effective_date(val: object) -> str | None:
    if isinstance(val, datetime):
        return val.date().strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return None
        try:
            parsed = datetime.fromisoformat(s)
            return parsed.date().strftime("%Y-%m-%d")
        except ValueError:
            pass
        try:
            parsed = date.fromisoformat(s)
            return parsed.strftime("%Y-%m-%d")
        except ValueError:
            pass
        if re.match(r"^\d{4}[-/]\d{2}[-/]\d{2}$", s):
            return s.replace("/", "-")
    return None


def _get_cell_value(ws, cell_ref: str) -> object | None:
    cell_obj = ws[cell_ref]
    if isinstance(cell_obj, tuple):
        if cell_obj and isinstance(cell_obj[0], tuple):
            cell_obj = cell_obj[0][0] if cell_obj[0] else None
        else:
            cell_obj = cell_obj[0] if cell_obj else None
    if cell_obj is None:
        return None
    return getattr(cell_obj, "value", None)


def inspect_excel_bytes_v2_op(inputs: dict[str, Any], store_) -> dict[str, Any]:
    input_profile_name = inputs.get("input_profile_name")
    if not input_profile_name:
        raise ValueError("input_profile_name is required")

    src_excel_bytes_key = inputs.get("src_excel_bytes_key")
    if not src_excel_bytes_key:
        raise ValueError("src_excel_bytes_key is required")

    src_excel_filename = inputs.get("src_excel_filename")
    if not src_excel_filename:
        raise ValueError("src_excel_filename is required")

    try:
        config = store_.configs.get_spec("input_profile", input_profile_name)
    except KeyError as e:
        raise ValueError(f"input_profile '{input_profile_name}' not found in configs") from e

    kind_rules = config.get("kind_rules")
    if kind_rules is None:
        raise ValueError(f"input_profile '{input_profile_name}' missing 'kind_rules' key")

    filename = Path(str(src_excel_filename)).name

    detected_kind = None
    matched_pattern = None
    for rule in kind_rules:
        pattern = rule.get("pattern")
        kind = rule.get("kind")
        if pattern and kind and re.match(pattern, filename):
            detected_kind = kind
            matched_pattern = pattern
            break

    date_rule = config.get("date_rule") or {}
    sheet = date_rule.get("sheet")
    cell = date_rule.get("cell")

    raw_date_value = None
    if sheet and cell:
        try:
            src_bytes = store_.get_bytes(src_excel_bytes_key)
            wb = openpyxl.load_workbook(BytesIO(src_bytes), data_only=True)
            if sheet in wb.sheetnames:
                ws = wb[sheet]
                raw_date_value = _get_cell_value(ws, cell)
        except Exception:
            raw_date_value = None

    effective_date = _normalize_effective_date(raw_date_value)

    result = {
        "schema_version": "inspect_result_v2",
        "input_profile_name": input_profile_name,
        "filename": filename,
        "detected_kind": detected_kind,
        "effective_date": effective_date,
        "evidence": {
            "matcher": "filename_regex",
            "matched_pattern": matched_pattern,
            "date": {
                "sheet": sheet,
                "cell": cell,
                "raw_value": raw_date_value,
            },
        },
    }

    return {"result": result}


def register(registry) -> None:
    registry.register("inspect", inspect_op)
    registry.register("inspect_filename_kind", inspect_filename_kind_op)
    registry.register(
        "inspect_excel_bytes_v2",
        inspect_excel_bytes_v2_op,
        required_inputs=(
            "input_profile_name",
            "src_excel_bytes_key",
            "src_excel_filename",
        ),
    )
