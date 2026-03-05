"""Fixture generation for hands-on / e2e tests."""

from __future__ import annotations

from pathlib import Path

# CSV fixture content for demo_csv_inspect / import_csv
DEMO_CSV_CONTENT = """LineNo,Item,Qty,Extra1,Extra2,Note,Date
1,a,10,x,100,memo1,2025-01-15
2,b,20,,200,,2024-12-01
3,c,30,z,300,memo3,2025-02-01
4,d,40,,,,"2024-06-01"
"""


def generate_csv(output_dir: str | Path) -> Path:
    """Generate demo_input.csv for demo_csv_inspect / import_csv. Returns output path."""
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "demo_input.csv"
    out_path.write_text(DEMO_CSV_CONTENT.strip(), encoding="utf-8")
    return out_path


def generate(output_dir: str | Path) -> Path:
    """Generate test_detect_region_input.xlsx. Returns output path."""
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "fixture generate requires openpyxl. Install with: pip install flowbook[excel]"
        ) from None

    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "test_detect_region_input.xlsx"

    wb = openpyxl.Workbook()
    meta = wb.active
    assert meta is not None
    meta.title = "meta"
    meta["A1"] = "Date"
    meta["B2"] = "2025-01-15"

    ws = wb.create_sheet("data", index=1)
    ws["A1"] = "Title"
    ws["A2"] = "Other"
    ws["E1"] = "Other table"
    ws["B6"] = "LineNo"
    ws["C6"] = "Item"
    ws["D6"] = "Qty"
    ws["E6"] = "Extra1"
    ws["F6"] = "Extra2"
    ws["G6"] = "Note"
    ws["H6"] = "Date"
    ws["B7"] = 1
    ws["C7"] = "a"
    ws["D7"] = 10
    ws["E7"] = "x"
    ws["F7"] = 100
    ws["G7"] = "memo1"
    ws["H7"] = "2025-01-15"
    ws["B8"] = 2
    ws["C8"] = "b"
    ws["D8"] = 20
    ws["E8"] = None
    ws["F8"] = 200
    ws["G8"] = None
    ws["H8"] = "2024-12-01"
    ws["B9"] = 3
    ws["C9"] = "c"
    ws["D9"] = 30
    ws["E9"] = "z"
    ws["F9"] = 300
    ws["G9"] = "memo3"
    ws["H9"] = "2025-02-01"
    ws["B10"] = 4
    ws["C10"] = "d"
    ws["D10"] = 40
    ws["G10"] = ""
    ws["H10"] = "2024-06-01"

    wb.save(out_path)
    return out_path


def generate_all(output_dir: str | Path) -> tuple[Path, Path]:
    """Generate both Excel and CSV fixtures. Returns (excel_path, csv_path)."""
    excel_path = generate(output_dir)
    csv_path = generate_csv(output_dir)
    return excel_path, csv_path
