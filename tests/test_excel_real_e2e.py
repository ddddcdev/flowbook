from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4

import openpyxl
import pytest

from extensions.steps.apply_mapping import ApplyMappingOp
from extensions.steps.inspect_excel_bytes_v2 import InspectExcelBytesV2Op
from extensions.steps.plan_from_template import PlanFromTemplateOp
from extensions.steps.read_excel_bytes import ReadExcelBytesOp
from extensions.steps.write_excel import WriteExcelOp
from flowbook.artifacts.memory_store import InMemoryArtifactsStore
from flowbook.configs.memory_store import InMemoryConfigStore
from flowbook.engine.engine import Engine
from flowbook.registry.extensions import register_steps
from flowbook.registry.registry import Registry

pytestmark = pytest.mark.e2e


def _resolve_template_name(config_store: InMemoryConfigStore, detected_kind: str | None) -> str:
    routing = config_store.get_spec("routing", "default")
    template_name = routing.get("map", {}).get(detected_kind, routing.get("default"))
    if template_name is None:
        raise RuntimeError(f"template_name is None for detected_kind={detected_kind}")
    return template_name


def test_excel_bytes_inspect_route_plan_execute_e2e() -> None:
    artifacts_store = InMemoryArtifactsStore()
    config_store = InMemoryConfigStore()

    registry = Registry()
    register_steps(registry)

    # ---- Configs ----
    input_profile_spec = {
        "kind_rules": [
            {"pattern": r"^fileA_.*\.xlsx$", "kind": "fileA"},
            {"pattern": r"^fileB_.*\.xlsx$", "kind": "fileB"},
        ],
        "date_rule": {"sheet": "meta", "cell": "B2"},
    }
    config_store.put_spec(
        kind="input_profile",
        name="source",
        spec=input_profile_spec,
        config_id=str(uuid4()),
    )

    routing_spec = {"map": {"fileA": "tmpl_fileA"}, "default": None}
    config_store.put_spec(
        kind="routing",
        name="default",
        spec=routing_spec,
        config_id=str(uuid4()),
    )

    mapping_spec = {
        "ops": [
            {"op": "select_cols", "cols": ["a", "b"]},
            {"op": "rename", "map": {"a": "A"}},
            {"op": "filter_rows", "expr": "A > 0"},
        ]
    }
    config_store.put_spec(
        kind="mapping",
        name="mvp_map",
        spec=mapping_spec,
        config_id=str(uuid4()),
    )

    template_spec = {
        "plan": {
            "steps": [
                {
                    "name": "read",
                    "op": "read_excel_bytes",
                    "inputs": {
                        ReadExcelBytesOp.Inputs.BYTES_KEY: "src_excel_bytes_key",
                        ReadExcelBytesOp.Inputs.SHEET: "sheet_name",
                        ReadExcelBytesOp.Inputs.HEADER: "header_row",
                        ReadExcelBytesOp.Inputs.OUT_KEY: "out_key_read",
                    },
                },
                {
                    "name": "map",
                    "op": "apply_mapping",
                    "inputs": {
                        ApplyMappingOp.Inputs.IN_KEY: "out_key_read",
                        ApplyMappingOp.Inputs.OUT_KEY: "out_key_map",
                        ApplyMappingOp.Inputs.MAPPING_NAME: "mapping_name_val",
                    },
                },
                {
                    "name": "write",
                    "op": "write_excel",
                    "inputs": {WriteExcelOp.Inputs.IN_KEY: "out_key_map"},
                },
            ]
        }
    }
    config_store.put_spec(
        kind="plan_template",
        name="tmpl_fileA",
        spec=template_spec,
        config_id=str(uuid4()),
    )

    engine = Engine(
        store=artifacts_store,
        registry=registry,
        config_store=config_store,
        meta={"env": "test"},
    )

    src_bytes = Path("tests/fixtures/excel/real_input.xlsx").read_bytes()
    bytes_artifact_key = "artifact:bytes/src_excel"

    # ---- Inspect (bytes + filename) ----
    inspect_run = engine.prepare()
    inspect_run.store.put_bytes(bytes_artifact_key, src_bytes)
    inspect_run.put_input("src_excel_bytes_key", bytes_artifact_key)
    inspect_run.put_input("src_excel_filename", "fileA_real_input.xlsx")
    inspect_run.put_input("input_profile_name", "source")

    inspect_config = {
        "steps": [
            {
                "name": "inspect",
                "op": "inspect_excel_bytes_v2",
                "inputs": {
                    InspectExcelBytesV2Op.Inputs.INPUT_PROFILE_NAME: "input_profile_name",
                    InspectExcelBytesV2Op.Inputs.SRC_EXCEL_BYTES_KEY: "src_excel_bytes_key",
                    InspectExcelBytesV2Op.Inputs.SRC_EXCEL_FILENAME: "src_excel_filename",
                },
            }
        ]
    }

    inspect_info = inspect_run.exec(pipeline_config=inspect_config)
    assert inspect_info.status == "succeeded", f"inspect failed: {inspect_info.errors}"

    result_key = inspect_info.steps[0].outputs[InspectExcelBytesV2Op.Outputs.RESULT]
    result = inspect_run.get_dict(result_key)
    assert result["detected_kind"] == "fileA"
    assert result["effective_date"] == "2026-02-10"

    template_name = _resolve_template_name(config_store, result["detected_kind"])

    # ---- Plan + Execute ----
    run = engine.prepare()
    run.store.put_bytes(bytes_artifact_key, src_bytes)
    run.put_input("src_excel_bytes_key", bytes_artifact_key)
    run.put_input("sheet_name", "data")
    run.put_input("header_row", 0)
    run.put_input("out_key_read", "artifact:df/in")
    run.put_input("mapping_name_val", "mvp_map")
    run.put_input("out_key_map", "artifact:df/mapped")
    run.put_input("template_name", template_name)

    planner_config = {
        "steps": [
            {
                "name": "planner",
                "op": "plan_from_template",
                "inputs": {PlanFromTemplateOp.Inputs.TEMPLATE_NAME: "template_name"},
            }
        ]
    }

    info1, info2 = run.exec_with_plan_once(planner_config=planner_config)

    assert info1.status == "succeeded", f"planner failed: {info1.errors}"
    assert info2.status == "succeeded", f"plan execution failed: {info2.errors}"

    write_step = info2.steps[-1]
    assert WriteExcelOp.Outputs.BYTES in write_step.outputs

    out_bytes_key = write_step.outputs[WriteExcelOp.Outputs.BYTES]
    out_bytes = run.get_bytes(out_bytes_key)
    assert len(out_bytes) > 0

    wb = openpyxl.load_workbook(BytesIO(out_bytes), data_only=True)
    assert wb.sheetnames
    assert "out" in wb.sheetnames
