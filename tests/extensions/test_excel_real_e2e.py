from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4

import openpyxl
import pytest

from flowbook import (
    Engine,
    InMemoryArtifactsStore,
    InMemoryConfigStore,
    Registry,
    register_steps,
)
from flowbook.core.configs.spec_types import (
    InputProfile,
    Mapping,
    PlanTemplate,
    Routing,
)
from flowbook.extensions.steps.apply_mapping import ApplyMappingOp
from flowbook.extensions.steps.inspect_excel_bytes_v2 import InspectExcelBytesV2Op
from flowbook.extensions.steps.plan_from_template import PlanFromTemplateOp
from flowbook.extensions.steps.read_excel_bytes import ReadExcelBytesOp
from flowbook.extensions.steps.write_excel import WriteExcelOp

pytestmark = pytest.mark.e2e


def _resolve_template_name(config_store: InMemoryConfigStore, detected_kind: str | None) -> str:
    routing = config_store.get_spec(Routing, "default")
    map_obj = routing.get("map") or {}
    template_name = (
        map_obj.get(detected_kind, routing.get("default"))
        if detected_kind is not None
        else routing.get("default")
    )
    if template_name is None:
        raise RuntimeError(f"template_name is None for detected_kind={detected_kind}")
    return template_name


def test_excel_bytes_inspect_route_plan_execute_e2e() -> None:
    artifacts_store = InMemoryArtifactsStore()
    config_store = InMemoryConfigStore()

    registry = Registry()
    register_steps(registry)

    # ---- Configs ----
    input_profile_spec = {
        "kind_rules": [
            {"pattern": r"^fileA_.*\.xlsx$", "kind": "fileA"},
            {"pattern": r"^fileB_.*\.xlsx$", "kind": "fileB"},
        ],
        "date_rule": {"sheet": "meta", "cell": "B2"},
    }
    config_store.put_spec(
        InputProfile,
        "source",
        input_profile_spec,
        config_id=str(uuid4()),
    )

    routing_spec = {"map": {"fileA": "tmpl_fileA"}, "default": None}
    config_store.put_spec(
        Routing,
        "default",
        routing_spec,
        config_id=str(uuid4()),
    )

    mapping_spec = {
        "ops": [
            {"op": "select_cols", "cols": ["a", "b"]},
            {"op": "rename", "map": {"a": "A"}},
            {"op": "filter_rows", "expr": "A > 0"},
        ]
    }
    config_store.put_spec(
        Mapping,
        "mvp_map",
        mapping_spec,
        config_id=str(uuid4()),
    )

    template_spec = {
        "plan": {
            "name": "excel_e2e",
            "steps": [
                {
                    "name": "read",
                    "op": "read_excel_bytes",
                    "inputs": {
                        ReadExcelBytesOp.Inputs.SRC_EXCEL_BYTES: "@src_excel_bytes",
                        ReadExcelBytesOp.Inputs.SHEET: "@sheet_name",
                        ReadExcelBytesOp.Inputs.HEADER: "@header_row",
                    },
                },
                {
                    "name": "map",
                    "op": "apply_mapping",
                    "inputs": {
                        ApplyMappingOp.Inputs.DF: "@read/df",
                        ApplyMappingOp.Inputs.MAPPING_NAME: "@mapping_name_val",
                    },
                },
                {
                    "name": "write",
                    "op": "write_excel",
                    "inputs": {WriteExcelOp.Inputs.DF: "@map/df"},
                },
            ]
        }
    }
    config_store.put_spec(
        PlanTemplate,
        "tmpl_fileA",
        template_spec,
        config_id=str(uuid4()),
    )

    engine = Engine(
        store=artifacts_store,
        registry=registry,
        config_store=config_store,
        meta={"env": "test"},
    )

    src_bytes = Path("tests/fixtures/excel/real_input.xlsx").read_bytes()
    bytes_artifact_key = "artifact/bytes/src_excel"

    # ---- Inspect (bytes + filename) ----
    inspect_run = engine.prepare()
    inspect_run.store.put_bytes(bytes_artifact_key, src_bytes)
    inspect_run.bind("src_excel_bytes", bytes_artifact_key)
    inspect_run.put_input("src_excel_filename", "fileA_real_input.xlsx")
    inspect_run.put_input("input_profile_name", "source")

    inspect_config = {
        "steps": [
            {
                "name": "inspect",
                "op": "inspect_excel_bytes_v2",
                "inputs": {
                    InspectExcelBytesV2Op.Inputs.INPUT_PROFILE_NAME: "@input_profile_name",
                    InspectExcelBytesV2Op.Inputs.SRC_EXCEL_BYTES: "@src_excel_bytes",
                    InspectExcelBytesV2Op.Inputs.SRC_EXCEL_FILENAME: "@src_excel_filename",
                },
            }
        ]
    }

    inspect_info = inspect_run.exec_plan(plan_config=inspect_config)
    assert inspect_info.status == "succeeded", f"inspect failed: {inspect_info.errors}"

    result_key = inspect_info.steps[0].outputs[InspectExcelBytesV2Op.Outputs.RESULT]
    result = inspect_run.get_dict(result_key)
    assert result["detected_kind"] == "fileA"
    assert result["effective_date"] == "2026-02-10"

    template_name = _resolve_template_name(config_store, result["detected_kind"])

    # ---- Plan + Execute ----
    run = engine.prepare()
    run.store.put_bytes(bytes_artifact_key, src_bytes)
    run.bind("src_excel_bytes", bytes_artifact_key)
    run.put_input("sheet_name", "data")
    run.put_input("header_row", 0)
    run.put_input("mapping_name_val", "mvp_map")
    run.put_input("template_name", template_name)

    planner_config = {
        "steps": [
            {
                "name": "planner",
                "op": "plan_from_template",
                "inputs": {PlanFromTemplateOp.Inputs.TEMPLATE_NAME: "@template_name"},
            }
        ]
    }

    info1, info2 = run.exec_with_planner_once(planner_config=planner_config)

    assert info1.status == "succeeded", f"planner failed: {info1.errors}"
    assert info2.status == "succeeded", f"plan execution failed: {info2.errors}"

    write_step = info2.steps[-1]
    assert WriteExcelOp.Outputs.BYTES in write_step.outputs

    out_bytes_key = write_step.outputs[WriteExcelOp.Outputs.BYTES]
    out_bytes = run.get_bytes(out_bytes_key)
    assert len(out_bytes) > 0

    wb = openpyxl.load_workbook(BytesIO(out_bytes), data_only=True)
    assert wb.sheetnames
    assert "out" in wb.sheetnames
