#!/usr/bin/env python3
"""Generate test_detect_region_input.xlsx (dirty layout: table not at A1, blank rows/cols).

Run from repo root. Overwrites tests/fixtures/excel/test_detect_region_input.xlsx.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl


def main() -> int:
    repo_root = Path(__file__).resolve().parent.parent
    out_path = repo_root / "tests" / "fixtures" / "excel" / "test_detect_region_input.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    meta = wb.active
    meta.title = "meta"
    meta["A1"] = "Date"
    meta["B2"] = "2025-01-15"

    ws = wb.create_sheet("data", index=1)
    ws["A1"] = "Title"
    ws["A2"] = "Other"
    ws["E1"] = "Other table"
    ws["B6"] = "明細No"
    ws["C6"] = "品名"
    ws["D6"] = "数量"
    ws["E6"] = "余計1"
    ws["F6"] = "余計2"
    ws["G6"] = "備考"
    ws["H6"] = "日付"
    ws["B7"] = 1
    ws["C7"] = "a"
    ws["D7"] = 10
    ws["E7"] = "x"
    ws["F7"] = 100
    ws["G7"] = "memo1"
    ws["H7"] = "2025-01-15"
    ws["B8"] = 2
    ws["C8"] = "b"
    ws["D8"] = 20
    ws["E8"] = None
    ws["F8"] = 200
    ws["G8"] = None
    ws["H8"] = "2024-12-01"
    ws["B9"] = 3
    ws["C9"] = "c"
    ws["D9"] = 30
    ws["E9"] = "z"
    ws["F9"] = 300
    ws["G9"] = "memo3"
    ws["H9"] = "2025-02-01"
    ws["B10"] = 4
    ws["C10"] = "d"
    ws["D10"] = 40
    ws["G10"] = ""
    ws["H10"] = "2024-06-01"

    wb.save(out_path)
    print(f"Wrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
